"""LISTA PEEK — responde la pregunta del founder: "dicen 'lista modificada'… ¿QUÉ se modificó?"

Cuando el vigía detecta que una LISTA DE PRECIOS cambió, baja SOLO ese archivo (Drive API,
$0, CERO IA) y lo lee con parsers deterministas: unidad por unidad, precio por precio.
Compara contra la FOTO ANTERIOR del contenido (vigia_listas_snapshot) — o contra el catálogo
de la plataforma la primera vez — y devuelve el diff humano: qué unidades cambiaron de precio,
cuáles ya no aparecen (¿vendidas?) y cuáles son nuevas. El resultado viaja pegado al evento,
al pendiente de la bandeja y a la tarjeta de Telegram.

Nota de doctrina: la ronda del vigía sigue siendo SOLO metadata. Esta es la única excepción
quirúrgica — un archivo, solo cuando su huella cambió, sin gastar un peso de IA.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

log = logging.getLogger("dmx.lista_peek")

_MAX_BYTES = 20 * 1024 * 1024          # una lista jamás pesa más; si pesa, no es lista
_PRECIO_MIN = 300_000                  # debajo de esto no es precio de depto (es m², fecha…)
_PRECIO_MAX = 200_000_000

_RE_PRECIO = re.compile(r"\$?\s*(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d{6,9}(?:\.\d+)?)")
# identificador de unidad al inicio de línea/celda: 'A-204', 'B 1402', '204', 'PH 2', 'GH1',
# 'LOCAL 3', 'Humbolt 201'… (letras opcionales + número, corto)
_RE_UNIDAD = re.compile(r"^([A-ZÁÉÍÓÚÑ]{0,12}[\s\-]?\d{1,4}[A-Z]?|PH\s?\d{0,3}|GH\s?\d{0,3}|"
                        r"LOCAL\s?\S{0,6}|ROOF\s?\S{0,10})\b", re.I)
_RE_STATUS = re.compile(r"vendid|apartad|reservad|bloquead|no\s+disponible|sold", re.I)


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def norm_unidad(u: Any) -> str:
    """'A-204' ≡ 'A 204' ≡ 'a204' — misma unidad, distinta pluma."""
    return re.sub(r"[\s\-–_.]+", "", str(u or "")).upper()


def _precio_de_linea(linea: str) -> Optional[float]:
    """El precio es el número MÁS GRANDE con pinta de dinero en la línea (las listas traen
    también m², niveles y fechas — todos más chicos que un precio de depto)."""
    vals = []
    for m in _RE_PRECIO.finditer(linea):
        try:
            v = float(m.group(1).replace(",", ""))
        except ValueError:
            continue
        if _PRECIO_MIN <= v <= _PRECIO_MAX:
            vals.append(v)
    return max(vals) if vals else None


def unidades_de_texto(texto: str) -> Dict[str, Dict[str, Any]]:
    """Texto plano (PDF/CSV) → {unidad_normalizada: {unidad, precio, status}}."""
    out: Dict[str, Dict[str, Any]] = {}
    for linea in (texto or "").splitlines():
        linea = linea.strip()
        if not linea:
            continue
        m = _RE_UNIDAD.match(linea)
        if not m:
            continue
        precio = _precio_de_linea(linea)
        status = "no_disponible" if _RE_STATUS.search(linea) else None
        if precio is None and status is None:
            continue                    # línea sin dato útil (encabezado, dirección…)
        uid = m.group(1).strip()
        k = norm_unidad(uid)
        if k and k not in out:          # la 1ª aparición manda (evita pies de página)
            out[k] = {"unidad": uid, "precio": precio, "status": status}
    return out


def _unidades_de_xlsx(data: bytes) -> Dict[str, Dict[str, Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: Dict[str, Dict[str, Any]] = {}
    for ws in wb.worksheets:
        for fila in ws.iter_rows(values_only=True):
            celdas = [c for c in fila if c is not None]
            if len(celdas) < 2:
                continue
            uid = None
            for c in celdas[:3]:        # la unidad vive en las primeras columnas
                s = str(c).strip()
                if _RE_UNIDAD.match(s) and len(s) <= 18:
                    uid = s
                    break
            if not uid:
                continue
            precio = None
            for c in celdas:
                try:
                    v = float(str(c).replace(",", "").replace("$", ""))
                except (TypeError, ValueError):
                    continue
                if _PRECIO_MIN <= v <= _PRECIO_MAX:
                    precio = max(precio or 0, v)
            texto_fila = " ".join(str(c) for c in celdas)
            status = "no_disponible" if _RE_STATUS.search(texto_fila) else None
            if precio is None and status is None:
                continue
            k = norm_unidad(uid)
            if k and k not in out:
                out[k] = {"unidad": uid, "precio": precio, "status": status}
    wb.close()
    return out


def _unidades_de_pdf(data: bytes) -> Dict[str, Dict[str, Any]]:
    import pdfplumber
    textos: List[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for pagina in pdf.pages[:30]:
            textos.append(pagina.extract_text() or "")
    return unidades_de_texto("\n".join(textos))


def leer_lista(nombre: str, mime: str, data: bytes) -> Dict[str, Dict[str, Any]]:
    """Bytes de la lista → unidades. Parser por formato; si el formato no se deja, {}."""
    n = (nombre or "").lower()
    m = (mime or "").lower()
    if "pdf" in m or n.endswith(".pdf"):
        return _unidades_de_pdf(data)
    if "sheet" in m or "excel" in m or n.endswith((".xlsx", ".xlsm")):
        return _unidades_de_xlsx(data)
    if "csv" in m or n.endswith(".csv"):
        return unidades_de_texto(data.decode("utf-8", errors="replace"))
    # último intento: como texto
    return unidades_de_texto(data.decode("utf-8", errors="replace"))


# ─── diff puro (testeable) ────────────────────────────────────────────────────
def diff_unidades(base: Dict[str, Dict[str, Any]], actual: Dict[str, Dict[str, Any]],
                  tope: int = 15) -> Dict[str, Any]:
    """base vs actual → el diff humano. Tope por lista para que Telegram no explote."""
    cambios_precio, ya_no_estan, nuevas, cambios_status = [], [], [], []
    for k, u in actual.items():
        b = base.get(k)
        if b is None:
            nuevas.append(u.get("unidad") or k)
            continue
        pa, pn = b.get("precio"), u.get("precio")
        if pa is not None and pn is not None and abs(pa - pn) > 1:
            cambios_precio.append({"unidad": u.get("unidad") or k, "antes": pa, "ahora": pn})
        if (b.get("status") or None) != (u.get("status") or None):
            cambios_status.append({"unidad": u.get("unidad") or k,
                                   "antes": b.get("status") or "disponible",
                                   "ahora": u.get("status") or "disponible"})
    for k, b in base.items():
        if k not in actual:
            ya_no_estan.append(b.get("unidad") or k)
    return {"cambios_precio": cambios_precio[:tope],
            "ya_no_estan": ya_no_estan[:tope], "nuevas": nuevas[:tope],
            "cambios_status": cambios_status[:tope],
            "totales": {"antes": len(base), "ahora": len(actual),
                        "cambios_precio": len(cambios_precio), "ya_no_estan": len(ya_no_estan),
                        "nuevas": len(nuevas), "cambios_status": len(cambios_status)}}


# ─── base de comparación desde el CATÁLOGO (la 1ª vez, sin snapshot previo) ───
async def _base_de_catalogo(db, fuente_id: str, dev_carpeta: str) -> Optional[Dict[str, Any]]:
    """Sin foto anterior de la lista, la referencia es lo que la plataforma ya publica del
    dev mapeado: se elige el desarrollo con más unidades en común con la lista leída."""
    m = await db.vigia_manifiesto.find_one(
        {"fuente_id": fuente_id, "dev_carpeta": dev_carpeta}, {"_id": 0, "dev_org_id": 1})
    org = (m or {}).get("dev_org_id")
    if not org:
        return None
    base_por_dev: Dict[str, Dict[str, Dict[str, Any]]] = {}
    async for d in db.developments.find({"developer_id": org}, {"_id": 0, "id": 1, "name": 1}):
        unidades: Dict[str, Dict[str, Any]] = {}
        async for u in db.units.find({"development_id": d["id"]},
                                     {"_id": 0, "unit_number": 1, "price": 1, "status": 1}):
            k = norm_unidad(u.get("unit_number"))
            if not k:
                continue
            st = (u.get("status") or "").lower()
            unidades[k] = {"unidad": u.get("unit_number"), "precio": u.get("price"),
                           "status": None if st in ("", "disponible", "available") else "no_disponible"}
        if unidades:
            base_por_dev[d["id"]] = {"nombre": d.get("name"), "unidades": unidades}
    return base_por_dev or None


def _mejor_base(base_por_dev: Dict[str, Dict[str, Any]],
                actual: Dict[str, Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    mejor, score = None, 0
    for _did, info in (base_por_dev or {}).items():
        n = len(set(info["unidades"]) & set(actual))
        if n > score:
            mejor, score = info, n
    return mejor if score >= 3 else None      # menos de 3 unidades en común = no es ese proyecto


# ─── el peek completo (descarga puntual + diff + snapshot) ────────────────────
async def peek_evento(db, conn: Dict[str, Any], fuente_id: str,
                      ev: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """Evento lista_cambiada/lista_nueva → {base, diff…} para el parte y la tarjeta.
    Siempre deja snapshot en vigia_listas_snapshot para que el PRÓXIMO cambio se compare
    lista-vs-lista (el diff perfecto). Fail-soft: None si el archivo no se deja leer."""
    import bulk_ingest_engine as bie
    a = ev.get("archivo") or {}
    if not conn or not a.get("id"):
        return None
    data, mime = await bie._download_file_bytes(conn, a["id"], a.get("mime") or "")
    if not data or len(data) > _MAX_BYTES:
        return None
    actual = leer_lista(a.get("nombre") or "", mime, data)
    if not actual:
        return {"base": None, "nota": "el formato de esta lista no se dejó leer sin IA",
                "totales": {"ahora": 0}}

    # FORENSE (07-17): renglones tapados/omitidos/entrelazados + filas sombreadas
    # (= apartadas, la marca que el texto no dice) — cada lista que cambia pasa por aquí
    forense: List[str] = []
    if "pdf" in (mime or "").lower() or (a.get("nombre") or "").lower().endswith(".pdf"):
        try:
            import lista_forense
            forense = lista_forense.alertas_forenses(data)
            for uid in lista_forense.unidades_sombreadas(data):
                k = norm_unidad(uid)
                if k in actual and not actual[k].get("status"):
                    actual[k]["status"] = "no_disponible"
                    actual[k]["status_fuente"] = "fila sombreada (apartada)"
        except Exception as e:  # noqa: BLE001 — el forense nunca tira el peek
            log.warning(f"[peek] forense: {e}")

    # base: la foto anterior de ESTA lista (huella distinta) — o el catálogo la 1ª vez
    snap_prev = await db.vigia_listas_snapshot.find_one(
        {"archivo_id": a["id"], "huella": {"$ne": a.get("huella")}},
        {"_id": 0}, sort=[("ts", -1)])
    if snap_prev:
        base, base_origen = snap_prev.get("unidades") or {}, "la versión anterior de la lista"
    else:
        por_dev = await _base_de_catalogo(db, fuente_id, ev.get("dev") or "")
        info = _mejor_base(por_dev or {}, actual)
        base = (info or {}).get("unidades") or {}
        base_origen = f"el catálogo de la plataforma ({(info or {}).get('nombre')})" if info else None
        # una lista puede cubrir SOLO una torre del desarrollo (Avalia: un PDF por torre) —
        # las unidades de las otras torres no "desaparecieron": se quitan de la base
        if base:
            prefijos = {re.match(r"^[A-Z]*", k).group(0) for k in actual}
            base = {k: v for k, v in base.items()
                    if re.match(r"^[A-Z]*", k).group(0) in prefijos}

    resultado = {"base": base_origen, "n_leidas": len(actual), **diff_unidades(base, actual)} if base else \
                {"base": None, "n_leidas": len(actual), "totales": {"ahora": len(actual)},
                 "nota": f"primera lectura: {len(actual)} unidades con precio — desde el "
                         f"próximo cambio te digo el diff exacto"}
    if forense:
        resultado["alertas_forenses"] = forense

    await db.vigia_listas_snapshot.update_one(
        {"archivo_id": a["id"], "huella": a.get("huella")},
        {"$set": {"archivo_id": a["id"], "huella": a.get("huella"),
                  "nombre": a.get("nombre"), "dev": ev.get("dev"),
                  "proyecto": ev.get("proyecto"), "ts": _now_iso(),
                  "unidades": actual, "n": len(actual)}}, upsert=True)
    return resultado


def _fmt_precio(v) -> str:
    try:
        return f"${float(v):,.0f}"
    except (TypeError, ValueError):
        return "$?"


def diff_plausible(cambios: Optional[Dict[str, Any]], n_catalogo: int) -> tuple:
    """GUARDIÁN DE PLAUSIBILIDAD (07-24): ¿el diff es CREÍBLE o huele a mal parseo? (ok, motivo).
    A diferencia de la Capa 0 del auditor (que mide VALOR por unidad), esto mide el TAMAÑO DEL CAMBIO
    y BLOQUEA el auto-apply ANTES de tocar el catálogo. Caso The Park: 79→153 con 141 'nuevas' y 67
    'desaparecidas' contra 153 en catálogo = lista re-numerada o mal leída, NO 141 altas + 67 ventas.
    Bloquear = NO aplicar solo, mandarlo a revisión (conservador: un falso positivo solo pide tu OK)."""
    t = (cambios or {}).get("totales") or {}
    n = max(int(n_catalogo or 0), int(t.get("antes") or 0), 1)
    nuevas = int(t.get("nuevas") or 0)
    fueron = int(t.get("ya_no_estan") or 0)
    tope = max(8, 0.4 * n)   # nadie agrega/pierde >40% (o >8) de su inventario de un jalón sin re-numerar
    if nuevas > tope:
        return False, (f"aparecen {nuevas} deptos 'nuevos' de golpe (tu catálogo tiene {n_catalogo}) — "
                       f"huele a lista re-numerada o mal leída, no a {nuevas} altas reales")
    if fueron > tope:
        return False, (f"desaparecen {fueron} deptos de golpe (tu catálogo tiene {n_catalogo}) — "
                       f"probable mal parseo del PDF, no {fueron} ventas de un jalón")
    return True, None


def es_cambio_real(cambios: Optional[Dict[str, Any]]) -> bool:
    """¿El peek trae un cambio de DATOS que amerita decisión del founder? (precio/estado/altas/bajas,
    o primera lectura de una lista nueva). False = re-subida IDÉNTICA (mismo archivo, mismos datos) →
    NO amerita tarjeta: el dev volvió a subir el PDF sin cambiar nada. Puro, testeable."""
    c = cambios or {}
    if c.get("cambios_precio") or c.get("cambios_status") or c.get("nuevas") or c.get("ya_no_estan"):
        return True
    if "primera lectura" in str(c.get("nota") or ""):
        return True
    return False


def forense_humano(cambios: Optional[Dict[str, Any]]) -> str:
    """Las alertas forenses (tokens crudos, 'caso Jai 25L'…) son para el DEV, no para el founder.
    Si el peek marcó algo raro en el PDF, esto lo dice en UNA línea humana, sin jerga ni basura."""
    if (cambios or {}).get("alertas_forenses"):
        return ("⚠️ La lista venía con el texto encimado o renglones tapados — puede faltar "
                "algún detalle; la reviso a fondo si me lo pides.")
    return ""


def lineas_de_cambios(cambios: Optional[Dict[str, Any]], sangria: str = "   ",
                      forense: bool = True) -> List[str]:
    """El diff → líneas humanas para el parte/tarjeta (puro, testeable).
    forense=False oculta los tokens crudos de la capa forense (los ve el dev, no el founder)."""
    if not cambios:
        return [f"{sangria}(no pude leer el detalle del archivo — se re-lee al aprobar)"]
    if cambios.get("nota"):
        return [f"{sangria}{cambios['nota']}"]
    t = cambios.get("totales") or {}
    out: List[str] = []
    for c in cambios.get("cambios_precio") or []:
        try:
            pct = (c["ahora"] - c["antes"]) / c["antes"] * 100
            out.append(f"{sangria}· unidad {c['unidad']}: {_fmt_precio(c['antes'])} → "
                       f"{_fmt_precio(c['ahora'])} ({pct:+.1f}%)")
        except (TypeError, ValueError, ZeroDivisionError, KeyError):
            out.append(f"{sangria}· unidad {c.get('unidad')}: cambió el precio")
    if t.get("cambios_precio", 0) > len(cambios.get("cambios_precio") or []):
        out.append(f"{sangria}… y {t['cambios_precio'] - len(cambios['cambios_precio'])} cambios de precio más")
    for c in cambios.get("cambios_status") or []:
        out.append(f"{sangria}· unidad {c['unidad']}: {c['antes']} → {c['ahora']}")
    if cambios.get("ya_no_estan"):
        us = ", ".join(str(x) for x in cambios["ya_no_estan"][:8])
        extra = t.get("ya_no_estan", 0) - min(len(cambios["ya_no_estan"]), 8)
        out.append(f"{sangria}· ya NO aparecen: {us}" + (f" (+{extra} más)" if extra > 0 else "")
                   + " — normalmente vendidas o apartadas")
    if cambios.get("nuevas"):
        us = ", ".join(str(x) for x in cambios["nuevas"][:8])
        out.append(f"{sangria}· nuevas en lista: {us}")
    if not out:
        out.append(f"{sangria}mismas unidades y precios que {cambios.get('base') or 'antes'} "
                   f"(cambió el archivo, no los datos)")
    elif cambios.get("base"):
        out.append(f"{sangria}(comparado contra {cambios['base']})")
    if forense:
        for alerta in cambios.get("alertas_forenses") or []:
            out.append(f"{sangria}{alerta}")
    return out
