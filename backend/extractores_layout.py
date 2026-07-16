"""EXTRACTORES POR FAMILIA DE LAYOUT — código puro donde antes iba IA.

El upgrade que cambia la economía de la ingesta (probado 07-15 con NUA/Nupol: 160/160
unidades, $0, instantáneo): cuando una lista pertenece a una FAMILIA conocida, se extrae
con parser determinista guiado por encabezados; la IA queda solo para layouts nunca
vistos — y cada layout nuevo se registra UNA vez. Familia #1: "vp_class" (las VP_Lista
de CLASS). Familia nueva = una entrada en FAMILIAS.

Doble validación integrada por familia (sus invariantes): habitables+exteriores=totales ·
crédito+enganche=precio. Lo que no valida NO pasa callado: sale en la pre-auditoría.
"""
from __future__ import annotations

import io
import re
import unicodedata
from typing import Any, Dict, List, Optional

from identidad_unidad import norm_unidad


def _num(s) -> Optional[float]:
    if s is None:
        return None
    t = re.sub(r"[^\d.]", "", str(s))
    try:
        return float(t) if t else None
    except ValueError:
        return None


def _norm_h(s) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", str(s or ""))
                if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().upper()


_CAMPOS_VP = {"DEPTO": "unidad", "BALCON": "m2_balcon", "TERRAZA": "m2_balcon",
              "PATIO": "m2_patio", "ROOF": "m2_roof", "BOD": "bodegas",
              "EST": "estacionamientos", "HABITABLE": "m2_habitable", "TOTAL": "m2_total",
              "PRECIO": "precio", "CREDITO": "credito", "ENGANCHE": "enganche",
              "RESERVACION": "reservacion", "CONTRATO": "contrato", "DIFERIR": "a_diferir"}
_POSICIONAL_VP = {0: "unidad", 1: "m2_balcon", 3: "m2_patio", 4: "bodegas",
                  5: "estacionamientos", 6: "m2_habitable", 7: "m2_total", 8: "precio",
                  9: "credito", 10: "enganche", 11: "reservacion"}


def _mapa_header(row) -> Optional[Dict[int, str]]:
    m: Dict[int, str] = {}
    for i, cell in enumerate(row):
        c = _norm_h(cell)
        for k, campo in _CAMPOS_VP.items():
            if k in c and campo not in m.values():
                m[i] = campo
                break
    return m if "unidad" in m.values() and "precio" in m.values() else None


def _resolver_m2(u: Dict[str, Any], celdas: List[str]) -> bool:
    """Par contiguo hab/tot + exterior CONFIRMADO por celda independiente (jamás derivado)."""
    nums = [_num(x) for x in celdas]
    for j in range(1, len(nums) - 1):
        a, b = nums[j], nums[j + 1]
        if not (a and b and 30 < a < 400 and a <= b < a + 90):
            continue
        ext = b - a
        if ext < 0.6:
            u["m2_habitable"], u["m2_total"] = a, b
            return True
        for k in range(1, j):
            if nums[k] is not None and abs(nums[k] - ext) < 0.06:
                u["m2_habitable"], u["m2_total"] = a, b
                u["m2_balcon"] = nums[k]
                if u.get("m2_patio") == nums[k]:
                    u["m2_patio"] = None
                return True
    return False


def extraer_vp(pdf_bytes: bytes) -> Dict[str, Any]:
    """La familia VP de CLASS: header-driven + fallback posicional + resolvedor por fila."""
    import pdfplumber
    unidades: Dict[str, Dict[str, Any]] = {}
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            t = page.extract_table() or []
            mapa = None
            for row in t:
                mh = _mapa_header(row)
                if mh:
                    mapa = mh
                    break
            for row in t:
                c = [str(x).strip() if x else "" for x in row]
                unidad = None
                for x in c[:3]:
                    x2 = x.replace("\n", "").replace(" ", "")
                    # formatos reales del catálogo CLASS: 203 · T2-1901 · PH01 (Dessea)
                    # · 02C (Jai) · G001 (Panorama) — el resto lo filtran las
                    # validaciones de m²/dinero de la propia familia
                    if re.match(r"^(T\d-?)?(PH\d{1,3}|[A-Z]?\d{3,4}|\d{2,4}[A-Z])$",
                                x2, re.IGNORECASE):
                        unidad = x.replace("\n", " ").strip()
                        break
                if not unidad:
                    continue
                usar = mapa or _POSICIONAL_VP
                u: Dict[str, Any] = {"unidad": unidad}
                for i, campo in usar.items():
                    if campo != "unidad" and i < len(c):
                        u[campo] = _num(c[i])
                hab, tot = u.get("m2_habitable"), u.get("m2_total")
                ext = (u.get("m2_balcon") or 0) + (u.get("m2_patio") or 0) + (u.get("m2_roof") or 0)
                ok_m2 = hab is not None and tot is not None and abs(hab + ext - tot) < 0.6
                if not ok_m2:
                    ok_m2 = _resolver_m2(u, c)
                # columna presente y vacía = CERO (la familia VP declara sus columnas)
                for cero in ("bodegas", "m2_patio", "m2_roof"):
                    if cero in [v for v in usar.values()] and u.get(cero) is None:
                        u[cero] = 0
                pr, cr, en = u.get("precio"), u.get("credito"), u.get("enganche")
                u["_valida_m2"] = bool(ok_m2)
                u["_valida_dinero"] = bool(pr and cr and en and abs((cr + en) - pr) < 2)
                unidades[norm_unidad(unidad)] = u
    us = list(unidades.values())
    return {"familia": "vp_class", "unidades": us,
            "validacion": {"m2": sum(1 for u in us if u["_valida_m2"]),
                           "dinero": sum(1 for u in us if u["_valida_dinero"]),
                           "total": len(us)}}


# ─── vp_texto: la MISMA familia VP leída por TEXTO (cuando la tabla se desalinea) ──
_MONEY_RE = re.compile(r"\$\s?[\d][\d\s,]*\.?\d*")
_UNIT_TOKEN = re.compile(
    r"^(T\d|[A-Z]{1,2}-?\d{1,4}[A-Z]?|\d{2,4}[A-Z]?|PH-?\d{0,3}|[A-Z]|\d{1,2}-)$",
    re.IGNORECASE)


def _tokens_de_unidad(tokens: List[str]) -> Optional[tuple]:
    """Cuántos tokens del inicio forman el número de unidad ('B 102' · '1- PH01' ·
    'T2 - 1901' · '0101' · 'G-001' · '02C'). None si la línea no arranca con unidad."""
    if not tokens or not _UNIT_TOKEN.match(tokens[0]):
        return None
    n, t0 = 1, tokens[0]
    if len(tokens) > 1:
        if t0.endswith("-") or (t0.isalpha() and len(t0) <= 2) or \
                re.fullmatch(r"T\d", t0, re.IGNORECASE):
            if tokens[1] == "-" and len(tokens) > 2:
                n = 3
            elif re.fullmatch(r"(PH)?\d{1,4}[A-Z]?", tokens[1], re.IGNORECASE):
                n = 2
    unidad = " ".join(tokens[:n]).replace(" - ", "-").strip()
    if not re.search(r"\d", unidad):
        return None                    # una unidad sin dígitos no es unidad
    return unidad, n


def extraer_vp_texto(pdf_bytes: bytes) -> Dict[str, Any]:
    """La familia VP por TEXTO crudo, posición-independiente (Revolución/Coyoacán/
    Dessea/Panorama/Jai la desalinean como tabla). Anatomía de la línea:
      unidad · [exteriores 0-3 decimales] · [conteos 0-2 enteros] · habitable ·
      total · $precio $crédito $enganche $reservación $contrato $a-diferir
    Asignación por INVARIANTE (no por posición): total = último decimal antes del
    dinero; habitable = penúltimo; el resto de decimales = exteriores en el orden
    del encabezado (balcón · patio · roof). 1 conteo = estacionamientos; 2 =
    bodegas + estacionamientos. Las validaciones son las mismas de la familia."""
    import pdfplumber
    unidades: Dict[str, Dict[str, Any]] = {}
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            for ln in (page.extract_text() or "").split("\n"):
                moneys = _MONEY_RE.findall(ln)
                if len(moneys) < 4:
                    continue                           # sin bloque de dinero no es unidad
                cabeza = ln[:ln.index(moneys[0])].strip()
                toks = cabeza.split()
                par = _tokens_de_unidad(toks)
                if not par:
                    continue
                unidad, n_u = par
                resto = toks[n_u:]
                numericos = [t for t in resto
                             if re.fullmatch(r"\d[\d,]*\.?\d*", t)]
                if len(numericos) < 2:
                    continue                           # sin hab+total no hay unidad
                # anatomía de la cola: [...ext/conteos...] + habitable + total —
                # los DOS últimos numéricos son hab/tot aunque vengan sin decimales
                # ('80', '128'); antes de ellos: enteros cortos = conteos (bod/est),
                # decimales = exteriores en el orden del encabezado
                u: Dict[str, Any] = {"unidad": unidad,
                                     "m2_total": _num(numericos[-1]),
                                     "m2_habitable": _num(numericos[-2])}
                previos = numericos[:-2]
                ext = [_num(t) for t in previos if re.fullmatch(r"\d[\d,]*\.\d+", t)]
                conteos = [int(t) for t in previos if re.fullmatch(r"\d{1,2}", t)]
                for campo, val in zip(("m2_balcon", "m2_patio", "m2_roof"), ext):
                    u[campo] = val
                if len(conteos) == 1:
                    u["estacionamientos"] = conteos[0]
                elif len(conteos) >= 2:
                    u["bodegas"], u["estacionamientos"] = conteos[0], conteos[1]
                dinero = [_num(m) for m in moneys]
                for campo, val in zip(("precio", "credito", "enganche", "reservacion",
                                       "contrato", "a_diferir"), dinero):
                    u[campo] = val
                hab, tot = u["m2_habitable"], u["m2_total"]
                u["_valida_m2"] = abs(hab + sum(ext) - tot) < 0.6
                pr, cr, en = u.get("precio"), u.get("credito"), u.get("enganche")
                u["_valida_dinero"] = bool(pr and cr and en and abs((cr + en) - pr) < 2)
                unidades[norm_unidad(unidad)] = u
    us = list(unidades.values())
    return {"familia": "vp_texto", "unidades": us,
            "validacion": {"m2": sum(1 for u in us if u["_valida_m2"]),
                           "dinero": sum(1 for u in us if u["_valida_dinero"]),
                           "total": len(us)}}


def extraer_vp_mejor(pdf_bytes: bytes) -> Dict[str, Any]:
    """La familia VP completa: intenta TABLA y TEXTO y se queda con la extracción
    que más unidades VÁLIDAS produce (la deriva de layout deja de doler)."""
    candidatos = []
    for fn in (extraer_vp, extraer_vp_texto):
        try:
            r = fn(pdf_bytes)
            v = r["validacion"]
            candidatos.append((min(v["m2"], v["dinero"]), v["total"], r))
        except Exception:  # noqa: BLE001
            continue
    if not candidatos:
        return {"familia": "vp_class", "unidades": [],
                "validacion": {"m2": 0, "dinero": 0, "total": 0}}
    candidatos.sort(key=lambda c: (c[0], c[1]), reverse=True)
    return candidatos[0][2]


def detecta_vp(nombre_archivo: str, primer_texto: str = "") -> bool:
    n = (nombre_archivo or "").lower()
    return "vp_lista" in n.replace(" ", "_") or "lista de precios" in (primer_texto or "").lower()[:200]


def extraer_maestro_class(xlsx_bytes: bytes) -> Dict[str, Any]:
    """Familia #2: el Inventario maestro de CLASS (Excel multi-dev, hoja CDMX).
    Devuelve unidades AGRUPADAS por DESARROLLO — el masivo toma su proyecto."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    unidades: List[Dict[str, Any]] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, r))
        if not d.get("PRODUCTO") or not d.get("DESARROLLO"):
            continue
        unidades.append({"desarrollo": str(d["DESARROLLO"]).strip(),
                         "_producto": str(d["PRODUCTO"]).strip(),
                         "unidad": norm_unidad(d["PRODUCTO"]),
                         "bedrooms": _num(d.get("RECAMARAS")),
                         "recamaras_opcion": _num(d.get("RECAMARAS OPCIONALES")),
                         "bathrooms": _num(d.get("BAÑOS")),
                         "estacionamientos": _num(d.get("ESTACIONAMIENTOS")),
                         "parking_type": str(d.get("TIPO DE ESTACIONAMIENTO") or "").strip().lower() or None,
                         "m2_habitable": _num(d.get("M2 HABITABLE")),
                         "m2_balcones": _num(d.get("M2 BALCONES / TERRAZA  / PATIO- JARDIN")),
                         "m2_roof": _num(d.get("M2 ROOF GARDEN")),
                         "m2_total": _num(d.get("M2 TOTAL")),
                         "precio": _num(d.get("PRECIO ACTUAL")),
                         "cuarto_servicio": str(d.get("CUARTO DE SERVICIO") or "").strip().lower() or None,
                         "amueblado": str(d.get("AMUEBLADO") or "").strip().lower() or None})
    return {"familia": "maestro_class", "unidades": unidades,
            "validacion": {"total": len(unidades)}}


def detecta_maestro(nombre_archivo: str, primer_texto: str = "") -> bool:
    n = (nombre_archivo or "").lower()
    return "inventario" in n and n.endswith((".xlsx", ".xls"))


UMBRAL_DRIFT = 0.90   # si una familia conocida valida <90%, el layout CAMBIÓ — alerta


def drift_de_familia(resultado: Dict[str, Any]) -> Optional[str]:
    """El detector de drift: la familia dejó de ajustar → alerta explícita, no fallo callado."""
    v = resultado.get("validacion") or {}
    total = v.get("total") or 0
    if not total or "m2" not in v:
        return None
    peor = min(v.get("m2", total), v.get("dinero", total)) / total
    if peor < UMBRAL_DRIFT:
        return (f"⚠ DRIFT de layout en familia '{resultado.get('familia')}': solo "
                f"{peor * 100:.0f}% valida — el dev cambió el formato; revisar antes de cargar")
    return None


# ─── FAMILIA GDC — lista header-driven (columnas bailan entre proyectos) ───────
# GDC no tiene Excel maestro (07-16): la lista PDF es la única fuente de unidad. Sus
# columnas cambian de orden/cantidad entre proyectos, así que se mapea por NOMBRE de
# encabezado, no por posición. El precio vive en columnas de esquema de pago
# ("10% 90%", "30% 20% 50%"…); el estatus va DENTRO de esa celda (VENDIDO/APARTADO/$).
_GDC_COL = {
    "NIVEL": "level", "DEPTO": "unit_number", "DEPTO.": "unit_number",
    "DEPARTAMENTO": "unit_number",
    "BALCON": "m2_balcony", "BALCÓN": "m2_balcony", "JARDIN/BALCON": "m2_balcony",
    "JARDÍN/BALCON": "m2_balcony",
    "TERRAZA": "m2_terrace", "ROOF O JARDIN": "m2_roof_garden", "ROOF": "m2_roof_garden",
    "ROOF/JARDIN": "m2_roof_garden", "ROOF GARDEN": "m2_roof_garden",
    "M2 TOTALES": "m2_total", "M2TOTALES": "m2_total", "M2": "m2_total",
    "CAJONES": "parking_spots",
    "RECAMARAS": "bedrooms", "RECÁMARAS": "bedrooms",
    "BANOS": "bathrooms", "BAÑOS": "bathrooms",
}
_GDC_ESQUEMA_RE = re.compile(r"^\s*\d{1,3}\s*%(\s+\d{1,3}\s*%)+\s*$")   # "10% 90%" · "30% 20% 50%"
# headers que NO son un dato faltante (vacíos, numeración, notas) → no cuentan como brecha
_GDC_IGNORA_H = {"", "#", "NO", "NO.", "ITEM", "CONSECUTIVO", "TIPO", "MODELO",
                 "PROTOTIPO", "OBSERVACIONES", "NOTAS", "STATUS", "ESTATUS", "ESTADO"}
_ESTATUS_TXT = {"VENDIDO": "vendido", "APARTADO": "reservado",
                "DISPONIBLE": "disponible", "BLOQUEADO": "bloqueado"}


def _dinero_gdc(s: str) -> Optional[float]:
    """'$ 1 4,310,848.82' → 14310848.82 (GDC mete espacios adentro del monto)."""
    t = re.sub(r"[^\d.]", "", str(s or ""))
    try:
        return float(t) if t else None
    except ValueError:
        return None


def extraer_gdc(pdf_bytes: bytes) -> Dict[str, Any]:
    """Familia GDC: header-driven. Mapea columnas por nombre, deriva interior =
    total − exteriores (M2 Totales incluye balcón/terraza, confirmado 07-16), lee el
    estatus de la celda de precio y guarda TODOS los esquemas de pago del encabezado."""
    import pdfplumber
    unidades: Dict[str, Dict[str, Any]] = {}
    no_mapeadas: set = set()
    campos_detectados: set = set()   # qué columnas del catálogo trae el encabezado
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            tbl = page.extract_table()
            if not tbl:
                continue
            # 1· hallar la fila de encabezado (trae DEPTO y M2)
            hdr_i, mapa, esquemas = None, {}, []
            for i, row in enumerate(tbl[:4]):
                cs = [_norm_h(c) for c in row]
                if any("DEPTO" in c or "DEPARTAMENTO" in c for c in cs) and \
                        any("M2" in c for c in cs):
                    hdr_i = i
                    for j, c in enumerate(cs):
                        if c in _GDC_COL:
                            mapa[j] = _GDC_COL[c]
                            campos_detectados.add(_GDC_COL[c])
                        elif _GDC_ESQUEMA_RE.match((row[j] or "").strip()):
                            esquemas.append({"col": j, "nombre": (row[j] or "").strip()})
                        elif c and c not in _GDC_IGNORA_H:
                            # columna con nombre que NI mapeamos NI es esquema → brecha de
                            # cobertura (capa 7): la fuente trae un dato que estamos tirando
                            no_mapeadas.add(c)
                    break
            if hdr_i is None or "unit_number" not in mapa.values():
                continue
            # 2· filas de datos
            for row in tbl[hdr_i + 1:]:
                c = [str(x).replace("\n", " ").strip() if x else "" for x in row]
                u: Dict[str, Any] = {}
                for j, campo in mapa.items():
                    if j < len(c) and c[j]:
                        v = c[j]
                        if campo in ("unit_number",):
                            u[campo] = v
                        elif campo == "level":
                            u[campo] = _num(v)
                        elif v.upper() not in ("NA", "N/A", "-"):
                            u[campo] = _num(v)
                num = str(u.get("unit_number") or "").strip()
                if not num or not re.search(r"\d", num):
                    continue
                es_local = "LOCAL" in num.upper()
                u["tipo"] = "local" if es_local else "departamento"
                # 3· esquemas de pago + estatus (la celda de precio manda)
                precios_scheme = []
                estatus = None
                for e in esquemas:
                    val = c[e["col"]] if e["col"] < len(c) else ""
                    up = val.upper().strip()
                    if any(k in up for k in _ESTATUS_TXT):
                        estatus = next(st for k, st in _ESTATUS_TXT.items() if k in up)
                    monto = _dinero_gdc(val)
                    if monto and monto > 10000:
                        precios_scheme.append({"esquema": e["nombre"], "precio": monto})
                if precios_scheme:
                    u["price_mxn"] = precios_scheme[0]["precio"]   # el 1er esquema = base
                    u["esquemas_pago"] = precios_scheme
                    u["status"] = "disponible"
                    # 10%/90% → enganche/crédito derivados del 1er esquema
                    pct = re.findall(r"\d{1,3}", precios_scheme[0]["esquema"])
                    if len(pct) == 2:
                        u["enganche_mxn"] = round(u["price_mxn"] * int(pct[0]) / 100)
                        u["credito_mxn"] = round(u["price_mxn"] * int(pct[1]) / 100)
                else:
                    u["status"] = estatus or "no_disponible"
                # 4· interior derivado (total incluye exteriores) + validación
                tot = u.get("m2_total")
                ext = sum(u.get(k) or 0 for k in ("m2_balcony", "m2_terrace", "m2_roof_garden"))
                if tot:
                    u["size_m2"] = round(tot - ext, 2) if tot - ext > 0 else tot
                u["_valida_m2"] = bool(tot and tot > 0)
                unidades[norm_unidad(num)] = u
    us = list(unidades.values())
    # campos CALCULADOS (no impresos en la lista) — el juez NO debe buscarlos en el PDF:
    # interior = total − exteriores · enganche/crédito = % del esquema de pago
    derivados = ["size_m2"]
    if any(u.get("enganche_mxn") for u in us):
        derivados += ["enganche_mxn", "credito_mxn"]
    return {"familia": "gdc", "unidades": us, "campos_derivados": derivados,
            "columnas_no_mapeadas": sorted(no_mapeadas),   # capa 7: fuente trae datos que no capturamos
            "campos_detectados": sorted(campos_detectados),  # columnas del catálogo en el header
            "cobertura_lista_ok": not no_mapeadas,         # ¿capturamos toda columna del header?
            "validacion": {"m2": sum(1 for u in us if u.get("_valida_m2")),
                           "dinero": sum(1 for u in us if u.get("price_mxn")),
                           "total": len(us),
                           "vendidas": sum(1 for u in us if u.get("status") == "vendido")}}


_BODEGA_RE = re.compile(
    r"^\s*\d+\s+([A-ZÁÉÍÓÚ ]+?)\s+([A-Z]-?\d+)\s+([\d.,]+)\s+(.+?)\s*$")


def extraer_gdc_bodegas(pdf_bytes: bytes) -> Dict[str, Any]:
    """Lista de BODEGAS de GDC (inventario aparte de los deptos): filas
    'NIVEL CLAVE M2 PRECIO' (SOTANO B-01 6.25 $260,000 · o APARTADO/VENDIDO en vez de $).
    Cada bodega es una unidad tipo='bodega' para que entre al mismo catálogo."""
    import pdfplumber
    us: List[Dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            for ln in (page.extract_text() or "").split("\n"):
                m = _BODEGA_RE.match(ln)
                if not m:
                    continue
                nivel, clave, m2, precio_txt = m.groups()
                if not re.search(r"\d", clave):
                    continue
                monto = _dinero_gdc(precio_txt)
                estatus = next((st for k, st in _ESTATUS_TXT.items()
                                if k in precio_txt.upper()), None)
                u: Dict[str, Any] = {
                    "unit_number": clave.strip(), "tipo": "bodega",
                    "nivel_texto": nivel.strip().title(),
                    "m2_total": _num(m2), "size_m2": _num(m2)}
                if monto and monto > 1000:
                    u["price_mxn"] = monto
                    u["status"] = "disponible"
                else:
                    u["status"] = estatus or "no_disponible"
                us.append(u)
    return {"familia": "gdc_bodegas", "unidades": us,
            "validacion": {"total": len(us),
                           "con_precio": sum(1 for u in us if u.get("price_mxn"))}}


def detecta_gdc(nombre_archivo: str, primer_texto: str = "") -> bool:
    n = (nombre_archivo or "").upper()
    return "_LP" in n or "ESQUEMA DE PAGO" in (primer_texto or "").upper()[:300]


# ─── EL REGISTRO (familia nueva = una entrada; la IA solo toca lo desconocido) ─
FAMILIAS: List[Dict[str, Any]] = [
    {"key": "vp_class", "detecta": detecta_vp, "extraer": extraer_vp_mejor,
     "nota": "VP_Lista de CLASS (tabla + texto posición-independiente, gana la "
             "extracción con más unidades válidas) · 160/160 (07-15) + masivo (07-15)"},
    {"key": "gdc", "detecta": detecta_gdc, "extraer": extraer_gdc,
     "nota": "Lista GDC header-driven (columnas variables · esquemas de pago · estatus "
             "en celda de precio · M2 Totales incluye exteriores) · 07-16"},
    {"key": "maestro_class", "detecta": detecta_maestro, "extraer": extraer_maestro_class,
     "nota": "Inventario maestro multi-dev de CLASS (Excel, hoja CDMX) · 582 renglones"},
]


def extraer_deterministico(nombre_archivo: str, pdf_bytes: bytes) -> Optional[Dict[str, Any]]:
    """Si la lista es de familia conocida → extracción con código puro ($0). None = IA."""
    for fam in FAMILIAS:
        try:
            if fam["detecta"](nombre_archivo):
                return fam["extraer"](pdf_bytes)
        except Exception:  # noqa: BLE001 — una familia rota no bloquea el fallback a IA
            continue
    return None
