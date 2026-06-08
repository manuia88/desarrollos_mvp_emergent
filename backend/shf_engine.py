"""
shf_engine — Índice SHF de Precios de la Vivienda (plusvalía OFICIAL · ING.3).
═══════════════════════════════════════════════════════════════════════════════
Hallazgo (reporte founder 2026-06-07): el índice SHF NO existe como serie en el SIE de Banxico
(cualquier ID de serie es inventado). Solo se publica como XLSX de datos abiertos en gob.mx
(trimestral · feb/may/ago/nov). Es la plusvalía/apreciación OFICIAL por región, nueva vs usada.

VERIFICADO (XLSX 1T2026 descargado · 10,285 filas · serie 2005–2026 trimestral):
  · Estructura: Consecutivo | Global | Estado | Municipio | Trimestre | Año | Indice.
  · CDMX tiene índice propio en SOLO 5 alcaldías (las de volumen hipotecario suficiente):
    Cuauhtémoc 185.20 (+4.7%) · Miguel Hidalgo 180.01 (+4.4%) · Benito Juárez 179.68 (+4.5%) ·
    Iztapalapa 173.12 (+5.2%) · Gustavo A. Madero 171.30 (+4.9%). CDMX estatal 175.20 (+4.5%).
  · CORRECCIÓN: el "+5.1%" del boletín es la ZM Valle de México (INCLUYE Edomex), NO la CDMX
    como entidad (+4.5%). Las otras 11 alcaldías heredan el factor estatal CDMX.
  · Base 2005 ≈ 33 (CDMX): el nivel ~175 implica que el precio nominal ~5x desde 2005.

Estrategia honesta (cero deuda · build for endstate):
  · SEED con los valores OFICIALES verificados 1T2026 (snapshot por alcaldía).
  · `ingest_series`: carga la serie trimestral completa (backend/data/shf_cdmx_serie.csv · 510 filas)
    a `shf_series` → alimenta la gráfica de plusvalía histórica.
  · `refresh_from_xlsx`: baja el XLSX oficial y parsea las filas de CDMX (snapshot + serie). Si no
    es alcanzable, conserva el seed (no inventa).
  · `get_appreciation(alcaldia)`: plusvalía anual oficial de la alcaldía (propia si es de las 5,
    si no hereda CDMX estatal) → ancla de plusvalía del modelo en la ficha de zona.
"""
from __future__ import annotations

import csv
import logging
import os
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

log = logging.getLogger("dmx.shf_engine")

# Snapshot OFICIAL verificado (Índice SHF · 1T2026 · gob.mx · XLSX descargado).
SHF_SEED: Dict[str, Any] = {
    "periodo": "2026-T1",
    "fuente": "Índice SHF de Precios de la Vivienda (datos abiertos · gob.mx)",
    "nacional_anual_pct": 8.7,
    "nueva_anual_pct": 9.1,
    "usada_anual_pct": 8.3,
    "avaluo_promedio": 2024337,
    "avaluo_mediana": 1331000,
    # CDMX como ENTIDAD (no la ZM Valle de México) — base de plusvalía de la ciudad.
    "cdmx_estatal": {"indice": 175.20, "plusvalia_anual_pct": 4.5},
    # Las 5 alcaldías con índice SHF propio (volumen hipotecario suficiente).
    "alcaldias": {
        "Cuauhtémoc": {"indice": 185.20, "plusvalia_anual_pct": 4.7},
        "Miguel Hidalgo": {"indice": 180.01, "plusvalia_anual_pct": 4.4},
        "Benito Juárez": {"indice": 179.68, "plusvalia_anual_pct": 4.5},
        "Iztapalapa": {"indice": 173.12, "plusvalia_anual_pct": 5.2},
        "Gustavo A. Madero": {"indice": 171.30, "plusvalia_anual_pct": 4.9},
    },
    # ZM Valle de México (INCLUYE Edomex) — métrica metropolitana, NO la CDMX como entidad.
    "zm_valle_mexico_anual_pct": 5.1,
    "url_xlsx": "https://www.gob.mx/cms/uploads/attachment/file/1077618/Indice_SHF_datos_abiertos_1_trim_2026.xlsx",
}

_SERIE_CSV = os.path.join(os.path.dirname(__file__), "data", "shf_cdmx_serie.csv")
_CDMX_ESTATAL = "CDMX (estatal)"

# Normaliza nombres de alcaldía para casar el campo `alcaldia` de las colonias con el SHF.
_ALC_ALIASES = {
    "cuauhtemoc": "Cuauhtémoc", "cuauhtémoc": "Cuauhtémoc",
    "miguel hidalgo": "Miguel Hidalgo",
    "benito juarez": "Benito Juárez", "benito juárez": "Benito Juárez",
    "iztapalapa": "Iztapalapa",
    "gustavo a. madero": "Gustavo A. Madero", "gustavo a madero": "Gustavo A. Madero",
    "gam": "Gustavo A. Madero",
}


def _iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _norm_alc(alcaldia: Optional[str]) -> Optional[str]:
    if not alcaldia:
        return None
    return _ALC_ALIASES.get(str(alcaldia).strip().lower())


async def ensure_shf(db) -> Dict[str, Any]:
    """Garantiza que el snapshot SHF esté en `shf_index` (siembra los valores oficiales si falta)."""
    try:
        doc = await db.shf_index.find_one({"_id": "current"}, {"_id": 0})
        if doc and doc.get("alcaldias"):
            return doc
        doc = {**SHF_SEED, "source": "seed_oficial", "updated_at": _iso()}
        await db.shf_index.update_one({"_id": "current"}, {"$set": doc}, upsert=True)
        return doc
    except Exception as e:
        log.warning(f"[shf] ensure: {e}")
        return dict(SHF_SEED)


async def get_appreciation(db, alcaldia: Optional[str] = None, region: Optional[str] = None) -> Dict[str, Any]:
    """Plusvalía anual OFICIAL (%) de la alcaldía. Si la alcaldía es una de las 5 con índice propio
    → usa el suyo; si no (o no se sabe) → hereda el factor estatal CDMX (+4.5%), marcado honesto.
    `region` se mantiene por compatibilidad (se ignora a favor de la alcaldía)."""
    doc = await ensure_shf(db)
    alc = _norm_alc(alcaldia)
    alcaldias = doc.get("alcaldias") or {}
    estatal = doc.get("cdmx_estatal") or {}
    if alc and alc in alcaldias:
        a = alcaldias[alc]
        return {
            "alcaldia": alc, "zona": alc, "es_propio": True,
            "plusvalia_anual_pct": a.get("plusvalia_anual_pct"), "indice": a.get("indice"),
            "periodo": doc.get("periodo"), "fuente": doc.get("fuente"), "es_oficial": True,
        }
    return {
        "alcaldia": alc, "zona": "Ciudad de México", "es_propio": False,
        "referencia": "Promedio CDMX (esta alcaldía no tiene índice propio del SHF)",
        "plusvalia_anual_pct": estatal.get("plusvalia_anual_pct"), "indice": estatal.get("indice"),
        "periodo": doc.get("periodo"), "fuente": doc.get("fuente"), "es_oficial": True,
    }


def _read_serie_csv() -> List[Dict[str, Any]]:
    """Lee la serie trimestral CDMX del CSV del repo → filas {alcaldia, anio, trimestre, indice}."""
    out: List[Dict[str, Any]] = []
    try:
        with open(_SERIE_CSV, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                try:
                    out.append({
                        "alcaldia": (r.get("Alcaldia") or "").strip(),
                        "anio": int(r["Anio"]), "trimestre": int(r["Trimestre"]),
                        "indice": float(r["Indice"]),
                    })
                except (TypeError, ValueError, KeyError):
                    continue
    except FileNotFoundError:
        log.warning(f"[shf] serie CSV no encontrado: {_SERIE_CSV}")
    return out


async def ingest_series(db, rows: Optional[List[Dict[str, Any]]] = None) -> Dict[str, Any]:
    """Carga la serie trimestral CDMX (2005–2026 · 6 series) en `shf_series`. Idempotente."""
    rows = rows if rows is not None else _read_serie_csv()
    if not rows:
        return {"ok": False, "reason": "Sin filas de serie (CSV ausente).", "cargadas": 0}
    n = 0
    for r in rows:
        if not r.get("alcaldia"):
            continue
        key = f"{r['alcaldia']}|{r['anio']}|{r['trimestre']}"
        await db.shf_series.update_one(
            {"_id": key},
            {"$set": {"alcaldia": r["alcaldia"], "anio": r["anio"],
                      "trimestre": r["trimestre"], "indice": r["indice"]}},
            upsert=True)
        n += 1
    return {"ok": True, "cargadas": n, "series": len(set(r["alcaldia"] for r in rows if r.get("alcaldia")))}


async def get_series(db, alcaldia: Optional[str] = None, desde_anio: int = 2018) -> Dict[str, Any]:
    """Serie de plusvalía (índice trimestral) para la gráfica de la ficha de zona. Usa la serie de
    la alcaldía si tiene propia, si no la estatal CDMX. Si `shf_series` está vacía, la siembra."""
    if await db.shf_series.count_documents({}) == 0:
        await ingest_series(db)
    alc = _norm_alc(alcaldia)
    target = alc if (alc and await db.shf_series.count_documents({"alcaldia": alc})) else _CDMX_ESTATAL
    cur = db.shf_series.find(
        {"alcaldia": target, "anio": {"$gte": desde_anio}},
        {"_id": 0, "anio": 1, "trimestre": 1, "indice": 1})
    pts = [p async for p in cur]
    pts.sort(key=lambda p: (p["anio"], p["trimestre"]))
    serie = [{"t": f"{p['anio']}-T{p['trimestre']}", "indice": p["indice"]} for p in pts]
    pv = None
    if len(serie) >= 5:  # plusvalía anual = índice vs mismo trimestre del año previo
        pv = round((serie[-1]["indice"] / serie[-5]["indice"] - 1) * 100, 1)
    return {"alcaldia": target, "es_propio": target != _CDMX_ESTATAL,
            "serie": serie, "plusvalia_anual_pct": pv, "fuente": SHF_SEED["fuente"]}


def _parse_xlsx_cdmx(content: bytes) -> Dict[str, Any]:
    """Parsea el XLSX oficial → snapshot por alcaldía CDMX + filas de serie. Estructura real:
    Consecutivo | Global | Estado | Municipio | Trimestre | Año | Indice."""
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)  # descarta encabezado
    serie: List[Dict[str, Any]] = []
    latest: Dict[str, Dict[int, float]] = {}  # alcaldia → {(anio*10+tri): indice} para snapshot
    snap_year = 0
    for row in rows:
        try:
            _cons, _glob, estado, muni, tri, anio, indice = row[:7]
        except (ValueError, TypeError):
            continue
        if not estado or "ciudad de m" not in str(estado).strip().lower():
            continue
        muni = (str(muni).strip() if muni else _CDMX_ESTATAL) or _CDMX_ESTATAL
        try:
            tri, anio, indice = int(tri), int(anio), float(indice)
        except (TypeError, ValueError):
            continue
        serie.append({"alcaldia": muni, "anio": anio, "trimestre": tri, "indice": indice})
        latest.setdefault(muni, {})[anio * 10 + tri] = indice
        snap_year = max(snap_year, anio)
    # snapshot: último trimestre disponible + plusvalía vs mismo trimestre del año previo
    alc_snapshot: Dict[str, Any] = {}
    estatal_snapshot = None
    for muni, pts in latest.items():
        if not pts:
            continue
        last_key = max(pts)
        idx = pts[last_key]
        prev = pts.get(last_key - 10)  # mismo trimestre, año previo
        pv = round((idx / prev - 1) * 100, 1) if prev else None
        snap = {"indice": round(idx, 2), "plusvalia_anual_pct": pv}
        if muni == _CDMX_ESTATAL:
            estatal_snapshot = snap
        else:
            alc_snapshot[muni] = snap
    return {"serie": serie, "alcaldias": alc_snapshot, "cdmx_estatal": estatal_snapshot, "anio": snap_year}


async def refresh_from_xlsx(db, url: Optional[str] = None) -> Dict[str, Any]:
    """Baja el XLSX oficial de SHF y refresca snapshot + serie de CDMX. Honesto: si no es alcanzable
    (CDN bloqueado / sin red), conserva el seed oficial y la serie del CSV del repo."""
    url = url or os.environ.get("IE_SHF_XLSX_URL") or SHF_SEED["url_xlsx"]
    try:
        import httpx
        async with httpx.AsyncClient(timeout=40, follow_redirects=True) as c:
            r = await c.get(url, headers={"User-Agent": "Mozilla/5.0 (DMX)"})
        data = r.content
        if r.status_code != 200 or data[:2] != b"PK" or len(data) < 10_000:
            await ingest_series(db)  # asegura la serie del repo aunque el CDN no responda
            cur = await ensure_shf(db)
            return {"ok": False, "reason": "XLSX no alcanzable desde aquí (CDN/redirect) — se conserva "
                    "el valor oficial sembrado + la serie del repo.", "current": cur}
        parsed = _parse_xlsx_cdmx(data)
        if parsed.get("alcaldias"):
            patch = {"alcaldias": {**SHF_SEED["alcaldias"], **parsed["alcaldias"]}, "updated_at": _iso()}
            if parsed.get("cdmx_estatal"):
                patch["cdmx_estatal"] = parsed["cdmx_estatal"]
            await db.shf_index.update_one({"_id": "current"}, {"$set": patch}, upsert=True)
        ser = await ingest_series(db, rows=parsed.get("serie") or None)
        return {"ok": True, "descargado_bytes": len(data),
                "alcaldias_actualizadas": len(parsed.get("alcaldias") or {}),
                "serie": ser}
    except Exception as e:
        log.warning(f"[shf] refresh: {e}")
        await ingest_series(db)
        cur = await ensure_shf(db)
        return {"ok": False, "reason": str(e), "current": cur}
