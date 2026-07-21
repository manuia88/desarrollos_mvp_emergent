"""Vigía de Quiero Casa (Google Sheet PÚBLICO) — corre HEADLESS, sin conector ni sesión.

Quiero Casa publica su portafolio en UN Google Sheet compartido "cualquiera con el enlace".
Su export público (`/export?format=xlsx`) responde 200 sin autenticación, así que el robot
horario lo baja solo. Cada pestaña de inventario = una 'lista' de su dev → se reusa
`lista_peek.diff_unidades` + `lista_apply.aplicar_cambios` (candado pelea, audit_log, alerta
de oportunidad), igual que las listas de CLASS/GDC en Drive. El baseline y el registro viven
en `vigia_listas_snapshot` (archivo_id `qc_sheet::<CODE>`) y `vigia_manifiesto`
(fuente_id `quiero_casa_sheet`).
"""
from __future__ import annotations

import io
import logging
import re
import unicodedata
import urllib.request
from collections import Counter
from datetime import datetime, timezone
from typing import Any, Dict

log = logging.getLogger("quiero_casa_sync")

FUENTE = "quiero_casa_sheet"
DEV_CARPETA = "QUIERO CASA (Google Sheets)"
SKIP_TABS = {"Dashboard", "_UpdateLog", "Resumen", "_ChartData"}
_STATUS = {"disponible": "disponible", "apartado": "apartado", "vendido": "vendido"}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _norm(s: Any) -> str:
    return unicodedata.normalize("NFKD", str(s or "").strip().lower()).encode("ascii", "ignore").decode()


def _num(v: Any):
    if v is None or str(v).strip().upper() in ("", "N/D", "ND", "-"):
        return None
    t = re.sub(r"[^0-9.\-]", "", str(v).replace(",", ""))
    try:
        return float(t) if t not in ("", "-", ".") else None
    except ValueError:
        return None


def _fetch_public_xlsx(sheet_id: str, timeout: int = 30) -> bytes:
    """Baja el Sheet PÚBLICO como XLSX (todas las pestañas). Sin auth."""
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (DMX-Vigia)"})
    with urllib.request.urlopen(req, timeout=timeout) as r:  # noqa: S310 (host fijo de Google)
        data = r.read()
    if data[:2] != b"PK":                      # PK = zip válido; si no, es login/HTML
        raise RuntimeError("el export no es un XLSX (¿la hoja dejó de ser pública?)")
    return data


def _codigo_estable(unidades: Dict[str, Dict[str, Any]]) -> str | None:
    """La identidad del proyecto vive en el SKU de los datos (IE1-0A-N001-P-0101 → 'IE1'),
    NO en el nombre de la pestaña (que el desarrollador edita: le agregó emojis 🔵). El prefijo
    de SKU dominante sobrevive a renombres de pestaña."""
    pref = [k.split("-")[0] for k in unidades if "-" in k and k.split("-")[0].isalnum()]
    return Counter(pref).most_common(1)[0][0] if pref else None


def _parse_tab(ws) -> Dict[str, Dict[str, Any]]:
    """Una pestaña de inventario → {SKU: {unidad, precio, status}} (formato de diff_unidades).
    Excluye comerciales (doctrina: catálogo residencial). Multi-torre → unidad con torre."""
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows)
               if r and "Torre" in [str(x).strip() for x in r]
               and "Estatus" in [str(x).strip() for x in r]), None)
    if hi is None:
        return {}
    ci = {_norm(h): j for j, h in enumerate(rows[hi])}

    def g(r, *names):
        for n in names:
            j = ci.get(_norm(n))
            if j is not None and j < len(r):
                return r[j]
        return None

    filas = [r for r in rows[hi + 1:]
             if r and str(g(r, "N° Depto", "N Depto") or "").strip() not in ("", "None")]
    torres = {str(g(r, "Torre") or "").strip() for r in filas if str(g(r, "Torre") or "").strip()}
    multi = len(torres) > 1
    out: Dict[str, Dict[str, Any]] = {}
    for r in filas:
        depto = str(g(r, "N° Depto", "N Depto") or "").strip()
        torre = str(g(r, "Torre") or "").strip() or None
        d = _norm(depto)
        if "local" in d or "oficina" in d or (d[:2] == "lc" and (len(d) == 2 or not d[2:3].isalpha())):
            continue                            # comercial: fuera
        sku = str(g(r, "SKU") or "").strip() or depto
        un = f"{torre}-{depto}" if (multi and torre) else depto
        est = _norm(g(r, "Estatus"))
        out[sku] = {"unidad": un, "precio": _num(g(r, "Precio")),
                    "status": _STATUS.get(est, "disponible" if est in ("", "none") else est)}
    return out


async def sincronizar(db, sheet_id: str | None = None) -> Dict[str, Any]:
    """Corrida headless: baja el Sheet público, diffea cada pestaña vs su baseline y aplica
    precio/estatus reusando lista_apply. Devuelve el resumen para el parte."""
    from openpyxl import load_workbook
    from lista_peek import diff_unidades
    from lista_apply import aplicar_cambios

    man = await db.vigia_manifiesto.find_one({"fuente_id": FUENTE}, {"_id": 0, "sheet_id": 1})
    sid = sheet_id or (man or {}).get("sheet_id")
    if not sid:
        return {"ok": False, "razon": "sin sheet_id en el manifiesto"}
    try:
        data = _fetch_public_xlsx(sid)
    except Exception as e:                       # noqa: BLE001 — fail-open
        log.warning(f"[quiero_casa] fetch falló: {e}")
        return {"ok": False, "razon": str(e)}

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    now = _now()
    tot = {"precios": 0, "estatus": 0, "peleas": 0, "senales_venta": 0,
           "proyectos_con_cambios": 0, "proyectos_revisados": 0}
    for s in wb.sheetnames:
        if s in SKIP_TABS:
            continue
        actual = _parse_tab(wb[s])
        if not actual:
            continue
        code = _codigo_estable(actual) or _norm(s.split(" ", 1)[-1])   # SKU manda; pestaña de respaldo
        tot["proyectos_revisados"] += 1
        archivo_id = f"qc_sheet::{code}"
        snap = await db.vigia_listas_snapshot.find_one({"archivo_id": archivo_id}, sort=[("ts", -1)])
        base = (snap or {}).get("unidades") or {}
        diff = diff_unidades(base, actual)
        t = diff["totales"]
        if t["cambios_precio"] or t["cambios_status"] or t["ya_no_estan"]:
            ev = {"dev": DEV_CARPETA,
                  "archivo": {"id": archivo_id, "nombre": f"Sheet Quiero Casa · {s}"},
                  "cambios": diff}
            try:
                res = await aplicar_cambios(db, FUENTE, ev)
            except Exception as e:               # noqa: BLE001
                log.warning(f"[quiero_casa] apply {code}: {e}")
                res = None
            if res:
                tot["precios"] += res.get("precios", 0)
                tot["estatus"] += res.get("status", 0)
                tot["peleas"] += res.get("peleas", 0)
                tot["senales_venta"] += res.get("senales_venta", 0)
                tot["proyectos_con_cambios"] += 1
            # registro hipersegmentado del evento (con el diff, no vacío)
            await db.vigia_eventos.insert_one({
                "id": f"vev_qc_{code}_{now[:10]}", "fuente_id": FUENTE, "ts": now,
                "tipo": "lista_cambiada", "dev": DEV_CARPETA, "proyecto": s,
                "detalle": {"archivo_id": archivo_id, "cambios": diff, "aplicado": res or {}}})
        # refrescar snapshot para el próximo diff
        await db.vigia_listas_snapshot.update_one(
            {"archivo_id": archivo_id},
            {"$set": {"archivo_id": archivo_id, "fuente_id": FUENTE, "unidades": actual,
                      "n": len(actual), "ts": now, "huella": f"sync_{now}"}}, upsert=True)
    tot["ok"] = True
    if tot["proyectos_con_cambios"]:
        log.info(f"[quiero_casa] {tot['proyectos_con_cambios']} proyecto(s) con cambios · "
                 f"{tot['precios']} precio · {tot['estatus']} estatus")
    return tot
